"""End-to-end QA audit harness for umn-elearning-assistant.

Run: venv/bin/python scripts/qa_audit.py
Produces PASS/FAIL per feature + writes data/metadata/qa_audit_report.json
"""
import importlib
import json
import sys
import tempfile
import traceback
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

RESULTS = []


def check(feature, name, fn):
    """Run fn -> (ok: bool, note: str). Records result, never raises."""
    try:
        ok, note = fn()
        RESULTS.append({"feature": feature, "test": name,
                        "status": "PASS" if ok else "FAIL", "note": str(note)})
    except Exception as e:
        RESULTS.append({"feature": feature, "test": name, "status": "FAIL",
                        "note": f"EXCEPTION {type(e).__name__}: {e}"})


def report():
    print("\n" + "=" * 78)
    print("QA AUDIT REPORT — umn-elearning-assistant")
    print("=" * 78)
    cur_feature = None
    n_pass = n_fail = 0
    for r in RESULTS:
        if r["feature"] != cur_feature:
            cur_feature = r["feature"]
            print(f"\n[{cur_feature}]")
        mark = "PASS" if r["status"] == "PASS" else "FAIL"
        if r["status"] == "PASS":
            n_pass += 1
        else:
            n_fail += 1
        print(f"  {mark}  {r['test']}")
        print(f"        -> {r['note']}")
    print("\n" + "-" * 78)
    print(f"TOTAL: {n_pass} PASS / {n_fail} FAIL / {len(RESULTS)} tests")
    print("-" * 78)
    out = ROOT / "data" / "metadata" / "qa_audit_report.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps({"summary": {"pass": n_pass, "fail": n_fail},
                               "results": RESULTS}, indent=2, ensure_ascii=False),
                   encoding="utf-8")
    print(f"JSON report: {out}")
    return n_fail


# ─────────────────────────────────────────────────────────────
# 1. MODULE & SYNTAX INTEGRITY
# ─────────────────────────────────────────────────────────────
MODULES = ["src.config", "src.moodle_client", "src.document_parser", "src.ai_service",
           "src.assignment_worker", "src.academic_styler", "src.conversation_manager",
           "src.anti_slop", "src.scheduler", "src.telegram_bot"]


def test_imports():
    for mod in MODULES:
        def _imp(mod=mod):
            try:
                m = importlib.import_module(mod)
                return True, f"imported {mod} ({Path(m.__file__).name})"
            except Exception as e:
                return False, f"{type(e).__name__}: {e}"
        check("1. Module Integrity", f"import {mod}", _imp)

    def _main():
        try:
            importlib.import_module("main")
            return True, "imported main.py (defines main())"
        except Exception as e:
            return False, f"{type(e).__name__}: {e}"
    check("1. Module Integrity", "import main.py", _main)

    def _compile():
        import py_compile
        bad = []
        for p in list((ROOT / "src").glob("*.py")) + [ROOT / "main.py", ROOT / "sync.py"]:
            try:
                py_compile.compile(str(p), doraise=True)
            except Exception as e:
                bad.append(f"{p.name}: {e}")
        return (not bad), ("all files compile cleanly" if not bad else "; ".join(bad))
    check("1. Module Integrity", "py_compile all sources", _compile)


# ─────────────────────────────────────────────────────────────
# 2. CONVERSATION MEMORY
# ─────────────────────────────────────────────────────────────
def test_conversation():
    from src.conversation_manager import ConversationManager, MAX_HISTORY_TURNS

    tmp = Path(tempfile.mkdtemp(prefix="qa_conv_"))
    mgr = ConversationManager(storage_dir=tmp)
    CHAT = 99001

    def _persist_dir():
        return True, f"storage dir auto-created: {tmp.exists()}"

    def _multi_turn():
        mgr.add_user_message(CHAT, "halo bot")
        mgr.add_assistant_message(CHAT, "halo juga")
        mgr.add_user_message(CHAT, "jelasin materi cyber dong")
        mgr.add_assistant_message(CHAT, "oke, CIA triad adalah ...")
        hist = mgr.get_history(CHAT)
        ok = len(hist) == 4 and hist[0]["role"] == "user" and hist[1]["role"] == "assistant"
        return ok, f"{len(hist)} turns, first={hist[0]['role']}, last={hist[-1]['role']}"
    check("2. Conversation Memory", "multi-turn add/get history", _multi_turn)

    def _persisted():
        f = tmp / f"history_{CHAT}.json"
        ok = f.exists()
        data = json.loads(f.read_text(encoding="utf-8")) if ok else []
        fresh = ConversationManager(storage_dir=tmp).get_history(CHAT)
        return ok and len(fresh) == 4, f"file={f.name} exists={ok}, reloaded={len(fresh)} turns"
    check("2. Conversation Memory", "JSON file persistence + reload", _persisted)

    def _trim():
        m = ConversationManager(storage_dir=tmp)
        for i in range(20):
            m.add_user_message(12345, f"msg {i}")
        hist = m.get_history(12345)
        return len(hist) == MAX_HISTORY_TURNS, f"kept {len(hist)}/{MAX_HISTORY_TURNS} (capped) last='{hist[-1]['text']}'"
    check("2. Conversation Memory", f"trim to MAX_HISTORY_TURNS", _trim)

    def _detect():
        def detect_fn(text):
            return "IF571" if "cyber" in text.lower() else None
        got = mgr.detect_active_course_from_history(CHAT, detect_fn)
        none_case = mgr.detect_active_course_from_history(999999, detect_fn)
        return got == "IF571" and none_case is None, f"detected={got!r}, empty-history={none_case!r}"
    check("2. Conversation Memory", "detect_active_course_from_history", _detect)

    def _format():
        s = mgr.format_history_for_prompt(CHAT)
        ok = "Mahasiswa: halo bot" in s and "AI Asisten: halo juga" in s
        return ok, f"{len(s)} chars, has both role labels: {ok}"
    check("2. Conversation Memory", "format_history_for_prompt", _format)

    def _clear():
        removed = mgr.clear_history(CHAT)
        gone = not (tmp / f"history_{CHAT}.json").exists()
        return removed and gone, f"clear_history returned {removed}, file gone={gone}"
    check("2. Conversation Memory", "clear_history", _clear)

    def _recent_summary():
        from src.config import ASSIGNMENT_OUTPUTS_FILE
        s = ConversationManager.get_recent_assignments_summary(max_items=3)
        raw = json.loads(ASSIGNMENT_OUTPUTS_FILE.read_text(encoding="utf-8")) if ASSIGNMENT_OUTPUTS_FILE.exists() else {}
        if not raw:
            return True, "SKIP-ish: assignment_outputs.json empty/Kosong -> returns ''.  (valid empty-case)"
        ok = "RIWAYAT TUGAS" in s and "Judul Tugas" in s
        return ok, f"source entries={len(raw)}, summary {len(s)} chars, header present={ok}"
    check("2. Conversation Memory", "get_recent_assignments_summary", _recent_summary)


# ─────────────────────────────────────────────────────────────
# 3. OFFICIAL UMN ACADEMIC STYLER
# ─────────────────────────────────────────────────────────────
def test_styler():
    from src.academic_styler import render_umn_academic_document
    from docx import Document
    from docx.shared import Cm

    tmp = Path(tempfile.mkdtemp(prefix="qa_styler_"))
    out = tmp / "Tugas_QA_Audit.docx"

    spec = {
        "filename": "Tugas QA Audit.docx",
        "sections": [
            {"heading": "BAB I PENDAHULUAN",
             "paragraphs": ["Paragraf uji audit QA untuk memverifikasi styler UMN.", "Paragraf kedua."],
             "bullets": ["Poin pertama", "Poin kedua"]},
            {"heading": "BAB II METODE",
             "paragraphs": ["Implementasi memakai Python."],
             "code": "def hello():\n    return 'halo UMN'\n"},
        ],
    }
    assignment = {"title": "Tugas QA Audit", "course_name": "(IF542-A) Deep Learning - LEC"}

    def _render():
        p = render_umn_academic_document(spec, assignment, "Kenny QA", "00000012345", out)
        return (p.exists() and p.stat().st_size > 0), f"{p.name} created, {p.stat().st_size if p.exists() else 0} bytes"
    check("3. UMN Styler", "render_umn_academic_document -> .docx valid", _render)

    doc = None
    try:
        doc = Document(str(out))
    except Exception as e:
        check("3. UMN Styler", "reopen .docx (validity)", lambda: (False, f"cannot open: {e}"))

    if doc is not None:
        def _margins():
            s = doc.sections[0]
            got = (round(s.top_margin.cm, 2), round(s.bottom_margin.cm, 2),
                   round(s.left_margin.cm, 2), round(s.right_margin.cm, 2))
            want = (4.0, 3.0, 4.0, 3.0)
            ok = got == want and round(s.page_width.cm, 1) == 21.0 and round(s.page_height.cm, 1) == 29.7
            return ok, f"margins T/B/L/R={got} (want {want}); A4={round(s.page_width.cm,1)}x{round(s.page_height.cm,1)}"
        check("3. UMN Styler", "margin 4-4-3-3 cm + A4 page", _margins)

        def _cover():
            txt = "\n".join(p.text for p in doc.paragraphs)
            need = ["UNIVERSITAS MULTIMEDIA NUSANTARA", "FAKULTAS TEKNIK DAN INFORMATIKA",
                    "PROGRAM STUDI INFORMATIKA", "Disusun Oleh", "NIM: 00000012345", "TANGERANG"]
            missing = [n for n in need if n not in txt]
            return not missing, ("cover contains all official elements" if not missing else f"missing: {missing}")
        check("3. UMN Styler", "cover resmi UMN (institusi/fakultas/NIM/Tangerang)", _cover)

        def _first_page_blank():
            s = doc.sections[0]
            return s.different_first_page_header_footer is True, \
                f"different_first_page_header_footer={s.different_first_page_header_footer} (cover bebas header/footer)"
        check("3. UMN Styler", "cover different first page header/footer", _first_page_blank)

        def _code_accent():
            if not doc.tables:
                return False, "no code-block table found"
            xml = doc.tables[0]._tbl.xml
            has_blue = "004C97" in xml
            has_bg = "F3F4F6" in xml
            mono = any(r.font.name == "Consolas" for p in doc.tables[0].cell(0, 0).paragraphs for r in p.runs)
            return (has_blue and has_bg and mono), \
                f"left border UMN-blue #004C97={has_blue}, bg #F3F4F6={has_bg}, Consolas font={mono}"
        check("3. UMN Styler", "code block blue UMN accent + gray bg", _code_accent)

        def _fonts():
            findings = set()
            for p in doc.paragraphs:
                for r in p.runs:
                    if r.font.name:
                        findings.add((r.font.name, round(r.font.size.pt, 1) if r.font.size else None))
            has_tnr12 = ("Times New Roman", 12.0) in findings
            return has_tnr12, f"font faces/sizes used: {sorted(findings)}"
        check("3. UMN Styler", "Times New Roman 12pt body font", _fonts)


# ─────────────────────────────────────────────────────────────
# 4. EXTERNAL CLOUD DOWNLOADER
# ─────────────────────────────────────────────────────────────
class FakeResp:
    def __init__(self, url="", status_code=200, text="", headers=None, content=b""):
        self.url = url
        self.status_code = status_code
        self.text = text
        self.headers = headers or {}
        self._content = content

    def raise_for_status(self):
        pass

    def iter_content(self, chunk_size=8192):
        yield self._content

    def json(self):
        return {}


def test_downloader():
    import src.moodle_client as mc

    def _with(target_url, title="Materi QA", ctype="application/pdf"):
        """Run _download_external_url against a mocked network, return (result, captured_url)."""
        tmp = Path(tempfile.mkdtemp(prefix="qa_dl_"))
        client = mc.MoodleClient()
        captured = {}

        # session.get handles first hop (moodle URL) and second hop (actual download)
        def fake_session_get(url, **kw):
            if "elearning.umn.ac.id/mod/url" in url:
                return FakeResp(url=target_url, status_code=200, text="<html>workaround</html>")
            captured["url"] = url
            return FakeResp(url=url, status_code=200, headers={"content-type": ctype},
                            content=b"%PDF-1.4 fake bytes")

        client.session.get = fake_session_get

        orig_get = mc.requests.get
        mc.requests.get = fake_session_get
        try:
            res = client._download_external_url("https://elearning.umn.ac.id/mod/url/view.php?id=1",
                                                title, tmp)
        finally:
            mc.requests.get = orig_get
        return res, captured.get("url"), tmp

    def _gdoc():
        res, url, tmp = _with("https://docs.google.com/document/d/ABC123xyz/edit")
        ok = url == "https://docs.google.com/document/d/ABC123xyz/export?format=pdf"
        file_ok = res and res["path"].endswith(".pdf") and Path(res["path"]).exists()
        return (ok and file_ok), f"export URL={url} (pdf export={ok}); file saved={bool(file_ok)} -> {res['path'] if res else None}"
    check("4. Cloud Downloader", "Google Docs -> PDF export rewrite", _gdoc)

    def _gsheet():
        res, url, tmp = _with("https://docs.google.com/spreadsheets/d/SHEET9/edit#gid=0")
        ok = url == "https://docs.google.com/spreadsheets/d/SHEET9/export?format=pdf"
        return ok, f"export URL={url}"
    check("4. Cloud Downloader", "Google Sheets -> PDF export rewrite", _gsheet)

    def _gslide():
        res, url, tmp = _with("https://docs.google.com/presentation/d/SLIDE7/edit")
        ok = url == "https://docs.google.com/presentation/d/SLIDE7/export?format=pdf"
        return ok, f"export URL={url}"
    check("4. Cloud Downloader", "Google Slides -> PDF export rewrite", _gslide)

    def _gdrive():
        res, url, tmp = _with("https://drive.google.com/file/d/DRV42/view?usp=sharing")
        ok = url == "https://drive.google.com/uc?export=download&id=DRV42"
        return ok, f"download URL={url}"
    check("4. Cloud Downloader", "Google Drive file -> uc?export=download", _gdrive)

    def _onedrive():
        res, url, tmp = _with("https://1drv.ms/u/s!AbCdEf123")
        ok = url is not None and "download=1" in url
        return ok, f"final URL={url} (download=1 appended={ok})"
    check("4. Cloud Downloader", "OneDrive 1drv.ms -> download=1", _onedrive)

    def _sharepoint():
        res, url, tmp = _with("https://umn.sharepoint.com/sites/x/Doc.aspx?sourcedoc=abc&file=ppt.pptx")
        ok = url is not None and "download=1" in url and "sourcedoc=abc" in url
        return ok, f"final URL={url}"
    check("4. Cloud Downloader", "SharePoint -> download=1 (preserve query)", _sharepoint)


# ─────────────────────────────────────────────────────────────
# 5. MOODLE SUBMISSION ENGINE — VALIDATION / PARAM HANDLING
# ─────────────────────────────────────────────────────────────
def test_submit_validation():
    import src.moodle_client as mc

    client = mc.MoodleClient()
    client.is_logged_in = True  # skip real login/network

    def _missing_file():
        res = client.submit_assignment("https://elearning.umn.ac.id/mod/assign/view.php?id=1",
                                       Path("/tmp/definitely_not_here_9f8a.docx"))
        ok = res.get("ok") is False and "tidak ditemukan" in res.get("error", "")
        return ok, f"ok={res.get('ok')}, error={res.get('error')!r}"
    check("5. Submission Engine", "file not found -> ok=False + clear error", _missing_file)

    def _empty_file():
        tmp = Path(tempfile.mkdtemp(prefix="qa_sub_")) / "empty.docx"
        tmp.write_bytes(b"")
        res = client.submit_assignment("https://elearning.umn.ac.id/mod/assign/view.php?id=1", tmp)
        ok = res.get("ok") is False and "kosong" in res.get("error", "")
        return ok, f"ok={res.get('ok')}, error={res.get('error')!r}"
    check("5. Submission Engine", "empty (0-byte) file rejected", _empty_file)

    def _http_error():
        c = mc.MoodleClient(); c.is_logged_in = True
        c.session.get = lambda url, **kw: FakeResp(url=url, status_code=404, text="")
        tmp = Path(tempfile.mkdtemp(prefix="qa_sub_")) / "f.docx"; tmp.write_bytes(b"data")
        res = c.submit_assignment("https://elearning.umn.ac.id/mod/assign/view.php?id=1", tmp)
        ok = res.get("ok") is False and "HTTP 404" in res.get("error", "")
        return ok, f"ok={res.get('ok')}, error={res.get('error')!r}"
    check("5. Submission Engine", "editsubmission HTTP!=200 handled", _http_error)

    def _no_form():
        c = mc.MoodleClient(); c.is_logged_in = True
        c.session.get = lambda url, **kw: FakeResp(url=url, status_code=200, text="<html><body>no form</body></html>")
        tmp = Path(tempfile.mkdtemp(prefix="qa_sub_")) / "f.docx"; tmp.write_bytes(b"data")
        res = c.submit_assignment("https://elearning.umn.ac.id/mod/assign/view.php?id=1", tmp)
        ok = res.get("ok") is False and "Form submission tidak ditemukan" in res.get("error", "")
        return ok, f"ok={res.get('ok')}, error={res.get('error')!r}"
    check("5. Submission Engine", "missing submission form -> clear error", _no_form)

    def _no_sesskey():
        c = mc.MoodleClient(); c.is_logged_in = True
        html = '<html><form id="mform1"><input name="foo" value="bar"/></form></html>'
        c.session.get = lambda url, **kw: FakeResp(url=url, status_code=200, text=html)
        tmp = Path(tempfile.mkdtemp(prefix="qa_sub_")) / "f.docx"; tmp.write_bytes(b"data")
        res = c.submit_assignment("https://elearning.umn.ac.id/mod/assign/view.php?id=1", tmp)
        ok = res.get("ok") is False and "Sesskey" in res.get("error", "")
        return ok, f"ok={res.get('ok')}, error={res.get('error')!r}"
    check("5. Submission Engine", "missing sesskey -> clear error", _no_sesskey)

    def _edit_url_param():
        # verifies ?action=editsubmission is appended correctly with & when URL already has query
        c = mc.MoodleClient(); c.is_logged_in = True
        seen = {}
        c.session.get = lambda url, **kw: (seen.__setitem__("url", url),
                                           FakeResp(url=url, status_code=404, text=""))[1]
        tmp = Path(tempfile.mkdtemp(prefix="qa_sub_")) / "f.docx"; tmp.write_bytes(b"data")
        c.submit_assignment("https://elearning.umn.ac.id/mod/assign/view.php?id=99", tmp)
        ok = seen.get("url", "").endswith("?id=99&action=editsubmission")
        return ok, f"GET url={seen.get('url')}"
    check("5. Submission Engine", "editsubmission query-string built correctly", _edit_url_param)


# ─────────────────────────────────────────────────────────────
# 6. COURSE ALIASES & RAG WEEK SCORING
# ─────────────────────────────────────────────────────────────
def test_aliases_and_rag():
    import src.ai_service as ai
    from src.ai_service import AIService

    svc = AIService()

    cases = [("stracom", "MSC5233"), ("stratcom", "MSC5233"), ("ai for stracom", "MSC5233"),
             ("cyber", "IF571"), ("cybersecurity", "IF571"),
             ("rti", "IF590"), ("english 3", "UM321"),
             ("deep learning", "IF542"), ("game dev", "IF581")]
    for q, want in cases:
        def _c(q=q, want=want):
            got = svc._detect_target_course(q)
            return got == want, f"'{q}' -> {got!r} (want {want})"
        check("6. Aliases & RAG", f"_detect_target_course('{q}')", _c)

    def _none():
        got = svc._detect_target_course("makan siang apa ya")
        return got is None, f"non-course query -> {got!r} (want None)"
    check("6. Aliases & RAG", "_detect_target_course(no-match) -> None", _none)

    def _weeknum():
        samples = {"Materi-IF571-M01-v261-01": 1, "Week 03_Perceptron": 3,
                   "Pertemuan 12 - Query": 12,
                   "Pertemuan ke 1_ Introduction to AI - An Overview": 1,
                   "RPKPS EM105": None}
        got = {k: AIService._module_week_number(k) for k in samples}
        ok = all(got[k] == v for k, v in samples.items())
        return ok, f"{got}"
    check("6. Aliases & RAG", "_module_week_number parsing", _weeknum)

    # RAG week scoring on a synthetic corpus (no touching real data)
    def _week_scoring():
        tmp = Path(tempfile.mkdtemp(prefix="qa_rag_"))
        cdir = tmp / "(IF542-A) Deep Learning - LEC"
        cdir.mkdir(parents=True)
        (cdir / "Materi-IF542-M01-Intro perceptron.txt").write_text(
            "perceptron " * 5 + "intro perceptron dasar", encoding="utf-8")
        (cdir / "Materi-IF542-M02-Backprop.txt").write_text(
            "backpropagation " * 3 + "deep learning jaringan", encoding="utf-8")

        orig = ai.EXTRACTED_TEXT_DIR
        ai.EXTRACTED_TEXT_DIR = tmp
        try:
            ctx = svc._get_relevant_context(query="deep learning pertemuan 2", max_chars=35000,
                                            target_code="IF542")
        finally:
            ai.EXTRACTED_TEXT_DIR = orig

        docs = [ln for ln in ctx.splitlines() if ln.startswith("=== DOKUMEN:")]
        first_is_m02 = bool(docs) and "M02" in docs[0]
        return first_is_m02, f"ranked docs={docs} (M02 first={first_is_m02})"
    check("6. Aliases & RAG", "week/pertemuan boost ranks right module first", _week_scoring)

    def _course_isolation():
        tmp = Path(tempfile.mkdtemp(prefix="qa_rag2_"))
        for name in ["(IF571-B) Cybersecurity - LEC", "(IF590-B) Information Technology Research - LEC"]:
            d = tmp / name; d.mkdir(parents=True)
            (d / "Materi Cyber.txt").write_text("cyber security cia triad " * 3, encoding="utf-8")
            (d / "Materi RTI.txt").write_text("metodologi penelitian skripsi " * 3, encoding="utf-8")
        orig = ai.EXTRACTED_TEXT_DIR
        ai.EXTRACTED_TEXT_DIR = tmp
        try:
            ctx = svc._get_relevant_context(query="cyber security", max_chars=35000, target_code="IF571")
        finally:
            ai.EXTRACTED_TEXT_DIR = orig
        docs = [ln for ln in ctx.splitlines() if ln.startswith("=== DOKUMEN:")]
        courses = [ln for ln in ctx.splitlines() if ln.startswith("=== MATA KULIAH:")]
        only_cyber = (not any("IF590" in ln for ln in courses)) and any("IF571" in ln for ln in courses)
        return only_cyber, f"courses={courses} docs={docs} (non-target IF590 excluded={only_cyber})"
    check("6. Aliases & RAG", "course isolation excludes non-target folders", _course_isolation)


# ─────────────────────────────────────────────────────────────
# 7. TELEGRAM BOT COMMANDS
# ─────────────────────────────────────────────────────────────
def test_telegram_commands():
    import src.telegram_bot as tb

    EXPECTED = {"/start", "/tugas", "/briefing", "/courses", "/kerjakan", "/kumpul",
                "/clear", "/reset", "/model", "/sync", "/id"}

    def _build():
        if not tb.TELEGRAM_BOT_TOKEN:
            tb.TELEGRAM_BOT_TOKEN = "123456:QA-FAKE-TOKEN"  # build() does not hit network
        app = tb.create_bot_app()
        if app is None:
            return False, "create_bot_app() returned None"
        return True, f"app built with {len(app.handlers.get(0, []))} handlers in group 0"
    check("7. Telegram Bot", "create_bot_app() builds", _build)

    def _commands():
        app = tb.create_bot_app()
        if app is None:
            return False, "no app"
        found = set()
        for hlist in app.handlers.values():
            for h in hlist:
                for c in getattr(h, "commands", set()) or set():
                    found.add("/" + c)
        missing = EXPECTED - found
        extra = found - EXPECTED
        return (not missing), f"registered={sorted(found)}; missing={sorted(missing)}; extra={sorted(extra)}"
    check("7. Telegram Bot", "all 11 command handlers registered", _commands)

    def _handlers_named():
        app = tb.create_bot_app()
        names = []
        for hlist in app.handlers.values():
            for h in hlist:
                cb = getattr(h, "callback", None)
                if cb:
                    names.append(cb.__name__)
        want = {"start_command", "tugas_command", "briefing_command", "courses_command",
                "kerjakan_command", "kumpul_command", "clear_command", "model_command",
                "sync_command", "id_command", "handle_message"}
        missing = want - set(names)
        return not missing, f"callbacks={sorted(set(names))}; missing={sorted(missing)}"
    check("7. Telegram Bot", "handler callbacks wired to right functions", _handlers_named)


def main():
    test_imports()
    test_conversation()
    test_styler()
    test_downloader()
    test_submit_validation()
    test_aliases_and_rag()
    test_telegram_commands()
    fails = report()
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
